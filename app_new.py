import os
import io
import json
import sqlite3
import jwt
import bcrypt
from datetime import datetime, timedelta
from typing import List, Optional
from fastapi import FastAPI, HTTPException, Depends, Header, File, UploadFile, Form
from fastapi.responses import FileResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from dotenv import load_dotenv
import uvicorn
from pathlib import Path

# Load environment variables
load_dotenv()
api_key = os.getenv("FIREWORKS_API_KEY")
jwt_secret = os.getenv("JWT_SECRET", "your-secret-key-change-this-in-production")
db_path = os.getenv("ADMIN_DB_PATH", "data/admin.db")
inactivity_timeout = int(os.getenv("ADMIN_INACTIVITY_TIMEOUT", "900"))

if not api_key:
    raise ValueError("FIREWORKS_API_KEY not found in .env file")

app = FastAPI()

# Create data directory if it doesn't exist
os.makedirs(os.path.dirname(db_path) or ".", exist_ok=True)

# --- Middleware Setup ---
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# --- RAG Globals ---
# Initialized at startup; kept in-memory for fast retrieval on every /chat call.
llm = None
embeddings = None
faiss_stores = {}           # {user_id (int): FAISS instance}
FAISS_BASE_DIR = "data/faiss_indexes"


def _load_all_faiss_stores():
    """Load all persisted FAISS indexes from disk into faiss_stores at startup."""
    global faiss_stores
    if embeddings is None:
        return
    try:
        from langchain_community.vectorstores import FAISS as _FAISS
        base = Path(FAISS_BASE_DIR)
        if not base.exists():
            return
        for user_dir in base.iterdir():
            if user_dir.is_dir() and (user_dir / "index.faiss").exists():
                try:
                    user_id = int(user_dir.name)
                    faiss_stores[user_id] = _FAISS.load_local(
                        str(user_dir), embeddings, allow_dangerous_deserialization=True
                    )
                    print(f"Loaded FAISS index for user_id={user_id}")
                except Exception as e:
                    print(f"Warning: could not load FAISS index at {user_dir}: {e}")
    except Exception as e:
        print(f"Warning: FAISS load failed: {e}")


@app.on_event("startup")
async def _startup_llm():
    """Initialize LLM, embeddings, and FAISS stores at application startup."""
    global llm, embeddings
    try:
        from langchain_fireworks import ChatFireworks as _ChatFireworks
        llm = _ChatFireworks(api_key=api_key, model="accounts/fireworks/models/kimi-k2p6")
    except Exception as e:
        print("Warning: could not initialize LLM on startup:", e)

    try:
        from langchain_fireworks import FireworksEmbeddings as _FWEmbed
        embeddings = _FWEmbed(
            api_key=api_key,
            model="accounts/fireworks/models/qwen3-embedding-8b"
        )
        print("Embeddings model initialized.")
    except Exception as e:
        print("Warning: could not initialize embeddings on startup:", e)

    _load_all_faiss_stores()


@app.on_event("shutdown")
async def _shutdown_llm():
    """Attempt to close/cleanup the LLM client on shutdown to avoid unclosed sessions."""
    global llm
    if not llm:
        return
    try:
        if hasattr(llm, "aclose") and callable(llm.aclose):
            await llm.aclose()
        elif hasattr(llm, "aclose_async") and callable(llm.aclose_async):
            await llm.aclose_async()
        elif hasattr(llm, "close") and callable(llm.close):
            llm.close()
    except Exception:
        pass
    finally:
        llm = None


# --- Database Setup ---
def init_db():
    """Initialize database with required tables"""
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS prompts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            description TEXT,
            content TEXT NOT NULL,
            is_default BOOLEAN DEFAULT FALSE,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS chat_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            name TEXT NOT NULL,
            inbound TEXT,
            reply TEXT,
            time TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
        )
    """)

    # Documents table for RAG pipeline
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS documents (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id     INTEGER NOT NULL,
            filename    TEXT NOT NULL,
            doc_type    TEXT NOT NULL,
            chunk_count INTEGER DEFAULT 0,
            vector_ids  TEXT,
            source_url  TEXT,
            is_active   INTEGER DEFAULT 1,
            created_at  TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
        )
    """)
    # Migrate existing DBs that lack the is_active column
    try:
        cursor.execute("ALTER TABLE documents ADD COLUMN is_active INTEGER DEFAULT 1")
        conn.commit()
    except Exception:
        pass  # Column already exists

    conn.commit()
    conn.close()


# Initialize DB on startup
init_db()


# --- Pydantic Models ---
class LoginRequest(BaseModel):
    username: str
    password: str

class RegisterRequest(BaseModel):
    username: str
    password: str

class PromptCreate(BaseModel):
    name: str
    description: str = ""
    content: str

class PromptUpdate(BaseModel):
    name: Optional[str] = None
    description: Optional[str] = None
    content: Optional[str] = None

class ChatHistoryEntry(BaseModel):
    name: str
    time: str
    inbound: str = ""
    reply: str = ""

class ChatRequest(BaseModel):
    message: str
    system_prompt: str = ""
    prompt_id: Optional[int] = None
    recent_messages: Optional[list] = None

class UrlDocRequest(BaseModel):
    url: str


# --- Helper Functions ---
def get_db():
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    return conn

def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode(), bcrypt.gensalt()).decode()

def verify_password(password: str, hash: str) -> bool:
    return bcrypt.checkpw(password.encode(), hash.encode())

def create_jwt_token(user_id: int, username: str, expires_in: int = 3600) -> str:
    payload = {
        "user_id": user_id,
        "username": username,
        "exp": datetime.utcnow() + timedelta(seconds=expires_in),
        "iat": datetime.utcnow()
    }
    return jwt.encode(payload, jwt_secret, algorithm="HS256")

def verify_jwt_token(token: str):
    try:
        payload = jwt.decode(token, jwt_secret, algorithms=["HS256"])
        return payload
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Invalid token")

def get_current_user(authorization: str = Header(None)):
    if not authorization:
        raise HTTPException(status_code=401, detail="Missing authorization header")
    scheme, token = authorization.split() if " " in authorization else (None, authorization)
    if scheme and scheme.lower() != "bearer":
        raise HTTPException(status_code=401, detail="Invalid authorization scheme")
    return verify_jwt_token(token)


# --- RAG: Text Extraction Helpers ---
def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract plain text from a PDF file using pypdf."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
        return "\n\n".join(pages)
    except Exception as e:
        raise ValueError(f"Could not read PDF: {e}")


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract plain text from a .docx file using python-docx."""
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)
    except Exception as e:
        raise ValueError(f"Could not read DOCX: {e}")


def extract_text_from_url(url: str) -> str:
    """Scrape and extract readable text from a web URL using requests + BeautifulSoup."""
    try:
        import requests as _requests
        from bs4 import BeautifulSoup
        headers = {"User-Agent": "Mozilla/5.0 (compatible; RAG-Bot/1.0)"}
        resp = _requests.get(url, headers=headers, timeout=15)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "html.parser")
        # Remove navigation, scripts, styles
        for tag in soup(["script", "style", "nav", "footer", "header", "aside"]):
            tag.decompose()
        text = soup.get_text(separator="\n")
        # Collapse blank lines
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        return "\n".join(lines)
    except Exception as e:
        raise ValueError(f"Could not fetch URL '{url}': {e}")


# --- RAG: Document Processing & Indexing ---
def process_and_index_document(
    user_id: int,
    text: str,
    filename: str,
    doc_type: str,
    source_url: str = ""
) -> List[str]:
    """
    Chunk text, create embeddings, add to the user's FAISS store,
    persist to disk, and return the list of FAISS vector IDs.
    """
    if not text or not text.strip():
        raise ValueError("Document contains no extractable text.")

    from langchain_text_splitters import RecursiveCharacterTextSplitter
    from langchain_core.documents import Document as LCDocument
    from langchain_community.vectorstores import FAISS as _FAISS
    from langchain_community.docstore.in_memory import InMemoryDocstore
    import faiss as _faiss
    from uuid import uuid4

    splitter = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=100)
    chunks = splitter.split_text(text)
    if not chunks:
        raise ValueError("Text produced no chunks after splitting.")

    docs = [
        LCDocument(
            page_content=chunk,
            metadata={"source": filename, "doc_type": doc_type, "url": source_url}
        )
        for chunk in chunks
    ]
    uuids = [str(uuid4()) for _ in docs]

    user_store = faiss_stores.get(user_id)
    if user_store is not None:
        # Incrementally add to existing index
        user_store.add_documents(documents=docs, ids=uuids)
    else:
        # Create a fresh FAISS index for this user
        dim = len(embeddings.embed_query("hello"))
        index = _faiss.IndexFlatL2(dim)
        user_store = _FAISS(
            embedding_function=embeddings,
            index=index,
            docstore=InMemoryDocstore(),
            index_to_docstore_id={}
        )
        user_store.add_documents(documents=docs, ids=uuids)
        faiss_stores[user_id] = user_store

    # Persist to disk
    index_dir = Path(FAISS_BASE_DIR) / str(user_id)
    index_dir.mkdir(parents=True, exist_ok=True)
    faiss_stores[user_id].save_local(str(index_dir))
    print(f"FAISS index saved: {index_dir} ({len(chunks)} chunks from '{filename}')")

    return uuids


# --- Auth Endpoints ---
@app.post("/admin/auth/login")
async def login(request: LoginRequest):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT id, username, password_hash FROM users WHERE username = ?", (request.username,))
    user = cursor.fetchone()
    conn.close()
    if not user or not verify_password(request.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Invalid credentials")
    token = create_jwt_token(user["id"], user["username"])
    return {"access_token": token, "token_type": "bearer", "user_id": user["id"], "username": user["username"]}


@app.post("/admin/auth/register")
async def register(request: RegisterRequest):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*) as count FROM users")
    user_count = cursor.fetchone()["count"]
    if user_count > 0:
        conn.close()
        raise HTTPException(status_code=403, detail="Admin already exists")
    password_hash = hash_password(request.password)
    try:
        cursor.execute(
            "INSERT INTO users (username, password_hash) VALUES (?, ?)",
            (request.username, password_hash)
        )
        conn.commit()
        user_id = cursor.lastrowid
        conn.close()
        token = create_jwt_token(user_id, request.username)
        return {"access_token": token, "token_type": "bearer", "user_id": user_id, "username": request.username}
    except Exception as e:
        conn.close()
        raise HTTPException(status_code=400, detail=str(e))


@app.get("/admin/auth/verify")
async def verify(current_user=Depends(get_current_user)):
    return {"valid": True, "user_id": current_user["user_id"], "username": current_user["username"]}


# --- Prompt Management Endpoints ---
@app.get("/admin/prompts")
async def list_prompts(current_user=Depends(get_current_user)):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT id, name, description, content, is_default, created_at, updated_at FROM prompts WHERE user_id = ? ORDER BY created_at DESC",
        (current_user["user_id"],)
    )
    prompts = [dict(row) for row in cursor.fetchall()]
    conn.close()
    return {"prompts": prompts}


@app.post("/admin/prompts")
async def create_prompt(prompt: PromptCreate, current_user=Depends(get_current_user)):
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute(
            "INSERT INTO prompts (user_id, name, description, content, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?)",
            (current_user["user_id"], prompt.name, prompt.description, prompt.content,
             datetime.utcnow().isoformat(), datetime.utcnow().isoformat())
        )
        conn.commit()
        prompt_id = cursor.lastrowid
        conn.close()
        return {"id": prompt_id, "message": "Prompt created successfully"}
    except Exception as e:
        conn.close()
        raise HTTPException(status_code=400, detail=str(e))


@app.put("/admin/prompts/{prompt_id}")
async def update_prompt(prompt_id: int, prompt: PromptUpdate, current_user=Depends(get_current_user)):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT id FROM prompts WHERE id = ? AND user_id = ?", (prompt_id, current_user["user_id"]))
    if not cursor.fetchone():
        conn.close()
        raise HTTPException(status_code=404, detail="Prompt not found")
    update_fields = []
    update_values = []
    if prompt.name:
        update_fields.append("name = ?")
        update_values.append(prompt.name)
    if prompt.description is not None:
        update_fields.append("description = ?")
        update_values.append(prompt.description)
    if prompt.content:
        update_fields.append("content = ?")
        update_values.append(prompt.content)
    update_fields.append("updated_at = ?")
    update_values.append(datetime.utcnow().isoformat())
    update_values.append(prompt_id)
    try:
        cursor.execute(f"UPDATE prompts SET {', '.join(update_fields)} WHERE id = ?", update_values)
        conn.commit()
        conn.close()
        return {"message": "Prompt updated successfully"}
    except Exception as e:
        conn.close()
        raise HTTPException(status_code=400, detail=str(e))


@app.delete("/admin/prompts/{prompt_id}")
async def delete_prompt(prompt_id: int, current_user=Depends(get_current_user)):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT id FROM prompts WHERE id = ? AND user_id = ?", (prompt_id, current_user["user_id"]))
    if not cursor.fetchone():
        conn.close()
        raise HTTPException(status_code=404, detail="Prompt not found")
    try:
        cursor.execute("DELETE FROM prompts WHERE id = ?", (prompt_id,))
        conn.commit()
        conn.close()
        return {"message": "Prompt deleted successfully"}
    except Exception as e:
        conn.close()
        raise HTTPException(status_code=400, detail=str(e))


@app.put("/admin/prompts/{prompt_id}/set-default")
async def set_default_prompt(prompt_id: int, current_user=Depends(get_current_user)):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT id FROM prompts WHERE id = ? AND user_id = ?", (prompt_id, current_user["user_id"]))
    if not cursor.fetchone():
        conn.close()
        raise HTTPException(status_code=404, detail="Prompt not found")
    try:
        cursor.execute("UPDATE prompts SET is_default = FALSE WHERE user_id = ?", (current_user["user_id"],))
        cursor.execute("UPDATE prompts SET is_default = TRUE WHERE id = ?", (prompt_id,))
        conn.commit()
        conn.close()
        return {"message": "Default prompt updated"}
    except Exception as e:
        conn.close()
        raise HTTPException(status_code=400, detail=str(e))


# --- Chat History Endpoints ---
@app.get("/admin/history")
async def get_chat_history(current_user=Depends(get_current_user)):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT id, name, inbound, reply, time FROM chat_history WHERE user_id = ? ORDER BY created_at DESC LIMIT 100",
        (current_user["user_id"],)
    )
    history = [dict(row) for row in cursor.fetchall()]
    conn.close()
    return {"history": history}


@app.post("/admin/history")
async def save_chat_history(entry: ChatHistoryEntry, current_user=Depends(get_current_user)):
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute(
            "INSERT INTO chat_history (user_id, name, inbound, reply, time) VALUES (?, ?, ?, ?, ?)",
            (current_user["user_id"], entry.name, entry.inbound, entry.reply, entry.time)
        )
        conn.commit()
        conn.close()
        return {"message": "Chat history saved"}
    except Exception as e:
        conn.close()
        raise HTTPException(status_code=400, detail=str(e))


@app.delete("/admin/history")
async def clear_history(current_user=Depends(get_current_user)):
    conn = get_db()
    cursor = conn.cursor()
    try:
        cursor.execute("DELETE FROM chat_history WHERE user_id = ?", (current_user["user_id"],))
        conn.commit()
        conn.close()
        return {"message": "Chat history cleared"}
    except Exception as e:
        conn.close()
        raise HTTPException(status_code=400, detail=str(e))


# --- Document Management Endpoints (RAG) ---
@app.get("/admin/documents")
async def list_documents(current_user=Depends(get_current_user)):
    """List all uploaded documents for the current user."""
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT id, filename, doc_type, chunk_count, source_url, is_active, created_at FROM documents WHERE user_id = ? ORDER BY created_at DESC",
        (current_user["user_id"],)
    )
    docs = [dict(row) for row in cursor.fetchall()]
    conn.close()
    return {"documents": docs}


@app.post("/admin/documents/upload")
async def upload_documents(
    files: List[UploadFile] = File(...),
    current_user=Depends(get_current_user)
):
    """Upload one or more documents (.txt, .pdf, .docx). Each is chunked and indexed in FAISS."""
    if embeddings is None:
        raise HTTPException(status_code=503, detail="Embedding model not ready. Please wait for server startup.")

    results = []
    conn = get_db()
    cursor = conn.cursor()

    for upload in files:
        filename = upload.filename or "unnamed"
        ext = Path(filename).suffix.lower().lstrip(".")

        if ext not in ("txt", "pdf", "docx"):
            results.append({"filename": filename, "status": "skipped", "reason": f"Unsupported type: .{ext}"})
            continue

        try:
            file_bytes = await upload.read()
            if ext == "pdf":
                text = extract_text_from_pdf(file_bytes)
                doc_type = "pdf"
            elif ext == "docx":
                text = extract_text_from_docx(file_bytes)
                doc_type = "docx"
            else:
                text = file_bytes.decode("utf-8", errors="replace")
                doc_type = "txt"

            vector_ids = process_and_index_document(
                user_id=current_user["user_id"],
                text=text,
                filename=filename,
                doc_type=doc_type
            )

            cursor.execute(
                "INSERT INTO documents (user_id, filename, doc_type, chunk_count, vector_ids, is_active) VALUES (?, ?, ?, ?, ?, 1)",
                (current_user["user_id"], filename, doc_type, len(vector_ids), json.dumps(vector_ids))
            )
            conn.commit()
            results.append({"filename": filename, "status": "ok", "chunks": len(vector_ids)})

        except Exception as e:
            results.append({"filename": filename, "status": "error", "reason": str(e)})

    conn.close()
    return {"results": results}


@app.post("/admin/documents/url")
async def add_url_document(request: UrlDocRequest, current_user=Depends(get_current_user)):
    """Scrape a URL, chunk and index its content in FAISS."""
    if embeddings is None:
        raise HTTPException(status_code=503, detail="Embedding model not ready. Please wait for server startup.")

    url = request.url.strip()
    if not url.startswith(("http://", "https://")):
        raise HTTPException(status_code=400, detail="URL must start with http:// or https://")

    try:
        text = extract_text_from_url(url)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    # Use the URL's hostname as the filename label
    from urllib.parse import urlparse
    hostname = urlparse(url).netloc or url[:40]
    filename = f"{hostname} (URL)"

    try:
        vector_ids = process_and_index_document(
            user_id=current_user["user_id"],
            text=text,
            filename=filename,
            doc_type="url",
            source_url=url
        )
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        "INSERT INTO documents (user_id, filename, doc_type, chunk_count, vector_ids, source_url, is_active) VALUES (?, ?, ?, ?, ?, ?, 1)",
        (current_user["user_id"], filename, "url", len(vector_ids), json.dumps(vector_ids), url)
    )
    conn.commit()
    doc_id = cursor.lastrowid
    conn.close()

    return {"id": doc_id, "filename": filename, "chunks": len(vector_ids), "message": "URL indexed successfully"}


@app.put("/admin/documents/{doc_id}/toggle")
async def toggle_document(doc_id: int, current_user=Depends(get_current_user)):
    """Flip the is_active flag for a document (enable/disable for RAG context)."""
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT id, is_active FROM documents WHERE id = ? AND user_id = ?",
        (doc_id, current_user["user_id"])
    )
    doc = cursor.fetchone()
    if not doc:
        conn.close()
        raise HTTPException(status_code=404, detail="Document not found")
    try:
        cursor.execute(
            "UPDATE documents SET is_active = 1 - is_active WHERE id = ? AND user_id = ?",
            (doc_id, current_user["user_id"])
        )
        conn.commit()
        new_state = 1 - doc["is_active"]
        conn.close()
        return {"id": doc_id, "is_active": bool(new_state)}
    except Exception as e:
        conn.close()
        raise HTTPException(status_code=400, detail=str(e))


@app.delete("/admin/documents/{doc_id}")
async def delete_document(doc_id: int, current_user=Depends(get_current_user)):
    """Delete a document record and remove its embeddings from the FAISS index."""
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT id, vector_ids FROM documents WHERE id = ? AND user_id = ?",
        (doc_id, current_user["user_id"])
    )
    doc = cursor.fetchone()
    if not doc:
        conn.close()
        raise HTTPException(status_code=404, detail="Document not found")

    # Remove embeddings from in-memory FAISS store
    user_id = current_user["user_id"]
    try:
        vector_ids = json.loads(doc["vector_ids"] or "[]")
        user_store = faiss_stores.get(user_id)
        if user_store and vector_ids:
            user_store.delete(ids=vector_ids)
            # Persist the updated index
            index_dir = Path(FAISS_BASE_DIR) / str(user_id)
            index_dir.mkdir(parents=True, exist_ok=True)
            user_store.save_local(str(index_dir))
    except Exception as e:
        print(f"Warning: could not remove vectors for doc {doc_id}: {e}")

    try:
        cursor.execute("DELETE FROM documents WHERE id = ?", (doc_id,))
        conn.commit()
        conn.close()
        return {"message": "Document deleted"}
    except Exception as e:
        conn.close()
        raise HTTPException(status_code=400, detail=str(e))


# --- Chat Endpoint ---
def extract_section(message_text: str, section_name: str) -> str:
    """Extract section from prompt message"""
    marker = f"{section_name}:"
    if marker not in message_text:
        return ""
    tail = message_text.split(marker, 1)[-1]
    for next_header in ["Latest thread context:", "Rules:", "Never use placeholders"]:
        if next_header in tail:
            tail = tail.split(next_header, 1)[0]
            break
    return tail.strip()


@app.post("/chat")
async def chat_endpoint(request: ChatRequest):
    """Chat endpoint with RAG context injection, admin prompt, and conversation history."""
    message_text = (request.message or "").strip()

    if not message_text:
        return {"reply": "Could you share the latest message text from the thread so I can respond accurately?"}

    latest_inbound = extract_section(message_text, "Latest inbound message")
    if len(latest_inbound) < 2:
        return {"reply": "Could you share the latest inbound message from the thread so I can reply accurately?"}

    DEFAULT_SYSTEM = (
        "You write LinkedIn direct-message replies on behalf of the account owner. "
        "You will be given a conversation thread and the latest inbound message. "
        "Read the full thread to understand context, but reply ONLY to the latest inbound message. "
        "Do not repeat anything already said in the thread. "
        "Output exactly one short reply in plain text (1-2 sentences, max 320 characters). "
        "Use the recipient first name naturally once if appropriate. "
        "NEVER invent, assume, or mention any identity, profession, skill, role, or company "
        "unless explicitly stated in this system prompt. "
        "If no identity context is provided, keep the reply neutral and conversational. "
        "Do not use placeholders or bracket variables. "
        "Do not produce templates, lists, headings, or generic introductions. "
        "If the latest inbound message is unclear, ask one short clarifying question."
    )

    # Determine system prompt (prioritize admin-configured prompt)
    active_system = DEFAULT_SYSTEM
    if request.system_prompt and request.system_prompt.strip():
        active_system = request.system_prompt.strip()
    elif request.prompt_id:
        try:
            conn = get_db()
            cursor = conn.cursor()
            cursor.execute("SELECT content FROM prompts WHERE id = ?", (request.prompt_id,))
            prompt_row = cursor.fetchone()
            conn.close()
            if prompt_row:
                active_system = prompt_row["content"]
                print(f"Using admin prompt (ID: {request.prompt_id})")
        except Exception:
            pass

    # Fall back to the admin's marked-default prompt from DB
    if active_system == DEFAULT_SYSTEM:
        try:
            conn = get_db()
            cursor = conn.cursor()
            cursor.execute("SELECT content FROM prompts WHERE is_default = 1 LIMIT 1")
            default_row = cursor.fetchone()
            conn.close()
            if default_row and default_row["content"]:
                active_system = default_row["content"]
                print("Using admin default prompt from DB")
        except Exception:
            pass

    # RAG: retrieve relevant context from active documents only
    rag_context = ""
    try:
        user_store = faiss_stores.get(1)  # single-admin system: always user_id=1
        if user_store:
            # Fetch the set of filenames the admin has enabled
            _conn = get_db()
            _cur = _conn.cursor()
            _cur.execute(
                "SELECT filename FROM documents WHERE user_id = 1 AND is_active = 1"
            )
            active_sources = {row["filename"] for row in _cur.fetchall()}
            _conn.close()

            if active_sources:
                # Retrieve a larger pool so filtering doesn't leave us empty-handed
                rag_docs_raw = user_store.similarity_search(latest_inbound, k=10)
                rag_docs = [
                    d for d in rag_docs_raw
                    if d.metadata.get("source") in active_sources
                ][:3]
                if rag_docs:
                    rag_context = "\n\n".join(d.page_content for d in rag_docs)
    except Exception as e:
        print(f"RAG retrieval warning: {e}")

    if rag_context:
        active_system = (
            f"{active_system}\n\n"
            f"--- Relevant context from your documents ---\n{rag_context}\n"
            f"--- End of context ---"
        )
        print(f"RAG: injected {len(rag_context)} chars of context")

    # Build the human message, optionally enriched with recent thread messages
    enhanced_message = message_text
    if request.recent_messages:
        messages_context = "Recent conversation:\n"
        for msg in request.recent_messages[-10:]:
            sender = msg.get("sender", "Unknown")
            text = msg.get("text", "")
            messages_context += f"{sender}: {text}\n"
        enhanced_message = f"{messages_context}\n{message_text}"

    response = llm.invoke([("system", active_system), ("human", enhanced_message)])
    reply_text = (response.content or "").strip()

    if not reply_text:
        return {"reply": "Thanks for your message. Could you share a little more context so I can respond clearly?"}

    if "[" in reply_text and "]" in reply_text:
        return {"reply": "Thanks for your message, and great to connect. Could you share a bit more detail on what you have in mind?"}

    return {"reply": reply_text}


# --- Static Files & Admin Portal ---
try:
    app.mount("/admin/static", StaticFiles(directory="linkedin_extension/admin"), name="static")
except Exception:
    pass

@app.get("/admin")
async def admin_index():
    try:
        return FileResponse("linkedin_extension/admin/index.html")
    except Exception:
        return {"message": "Admin portal not found."}

@app.get("/admin/auth")
async def admin_auth():
    try:
        return FileResponse("linkedin_extension/admin/auth.html")
    except Exception:
        return {"message": "Auth page not found"}

@app.get("/admin/{path:path}")
async def admin_files(path: str):
    """Fallback route to serve admin static files under linkedin_extension/admin."""
    base = Path("linkedin_extension") / "admin"
    if not path or path in ("/", "index.html"):
        index = base / "index.html"
        if index.exists():
            return FileResponse(str(index))
    candidate = base / path
    if candidate.exists() and candidate.is_file():
        return FileResponse(str(candidate))
    if path == "favicon.ico":
        fav = base / "favicon.ico"
        if fav.exists():
            return FileResponse(str(fav))
    raise HTTPException(status_code=404, detail="Not found")


if __name__ == "__main__":
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT COUNT(*) as count FROM users")
    user_count = cursor.fetchone()["count"]
    conn.close()
    if user_count == 0:
        print("\n" + "=" * 60)
        print("NO ADMIN FOUND!")
        print("Run: python setup_admin.py")
        print("to create your first admin account")
        print("=" * 60 + "\n")
    uvicorn.run("app_new:app", host="0.0.0.0", port=9011, reload=True)
